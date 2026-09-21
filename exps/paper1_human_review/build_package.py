"""Build blinded, offline annotation deliveries without changing frozen samples."""
import csv
import hashlib
import html
import json
import sys
import zipfile
from pathlib import Path
import markdown

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from exps.paper1_mechanism_audit.protocol import read_jsonl

HERE = Path(__file__).resolve().parent
OUT = HERE / 'delivery'
AUDIT = ROOT / 'exps/paper1_mechanism_audit/human_review'
EXT = ROOT / 'exps/paper1_submission_extensions'
COLUMNS = ['task', 'item_id', 'source', 'schema', 'input_graph', 'output_graph', 'operation', 'triple', 'is_error', 'repair_acceptable', 'notes']


def read_csv(path):
    with path.open(newline='', encoding='utf-8') as f:
        return list(csv.DictReader(f))


def items():
    cases = {x['case_id']: x for x in read_jsonl(ROOT/'exps/paper1_repair_benchmark/benchmark.jsonl')}
    extracted = {x['case_id']: x for x in read_jsonl(EXT/'natural_extraction_claude.jsonl')}
    mappings = {x['item_id']: x for x in read_csv(AUDIT/'coordinator_mapping.csv')}
    result, private = [], []
    for row in read_csv(EXT/'natural_error_annotation_sample.csv'):
        cid = row['case_id']
        operation = 'add' if row['issue_type'] == 'missing_gold_triple' else 'remove'
        triple = {k: row[k] for k in ['head', 'relation', 'tail']}
        result.append({'task': 'D', 'item_id': row['sample_id'], 'source': row['source_evidence'],
            'schema': cases[cid]['allowed_relations'], 'input_graph': extracted[cid]['triples'],
            'output_graph': None, 'operation': operation, 'triple': triple})
        private.append({'task':'D','item_id':row['sample_id'],'case_id':cid,'configuration':'input_discrepancy','operation':operation})
    for row in read_csv(AUDIT/'annotator_a.csv'):
        mapping = mappings[row['item_id']]; cid = mapping['case_id']
        result.append({'task': 'E', **{k: row[k] for k in ['item_id','source','operation']},
            'schema': cases[cid]['allowed_relations'],
            **{k: json.loads(row[k]) for k in ['input_graph','output_graph','triple']}})
        private.append({'task':'E','item_id':row['item_id'],'case_id':cid,'configuration':mapping['method'],'operation':row['operation']})
    assert len(result) == 400
    return result, private


def workbook(records, path, who):
    wb = Workbook(); intro = wb.active; intro.title = '使用说明'
    instructions = [f'标注者 {who}：400 条，两个任务分别统计',
        '请先阅读随包的 ANNOTATOR_GUIDE_zh.md，优先使用离线 HTML 查看长文本。',
        '只填写黄色三列：is_error、repair_acceptable、notes。标签为 1、0、U。',
        'D：输入差异；E：实际修改。D 的 output_graph 留空，修改仅为建议。',
        '需要独立标注；不要查看他人标签、论文结果或协调员映射。',
        '保存最终 XLSX 交回协调员。不要修改 ID、原文、schema 或图谱。']
    for x in instructions: intro.append([x])
    intro.column_dimensions['A'].width=110
    for row in intro:
        row[0].alignment=Alignment(wrap_text=True,vertical='top')
    for task,title in [('D','输入差异'),('E','实际修改')]:
        ws=wb.create_sheet(title);ws.append(COLUMNS)
        for item in records:
            if item['task']!=task:continue
            ws.append([json.dumps(item[k],ensure_ascii=False,indent=1) if isinstance(item.get(k),(dict,list))
                       else item.get(k,'') for k in COLUMNS])
        ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
        widths=[7,15,65,24,58,58,12,42,15,23,48]
        from openpyxl.utils import get_column_letter
        for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
        dv=DataValidation(type='list',formula1='"1,0,U"',allow_blank=True)
        dv.error='请填写 1、0 或 U';dv.showErrorMessage=True;ws.add_data_validation(dv);dv.add('I2:J201')
        for row in ws:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(wrap_text=True,vertical='top')
                cell.font=Font(name='Microsoft YaHei',size=10)
                if cell.row==1:
                    cell.fill=PatternFill('solid',fgColor='DCE6F1');cell.font=Font(name='Microsoft YaHei',bold=True)
                elif cell.column>=9:
                    cell.fill=PatternFill('solid',fgColor='FFF2CC')
        for i in range(2,202):ws.row_dimensions[i].height=120
    wb.save(path)


HTML = r'''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Paper 1 独立标注 __WHO__</title>
<style>body{font:17px/1.65 system-ui,sans-serif;margin:0;background:#f4f6f9;color:#172536}header{background:#12334c;color:white;padding:14px 4%;position:sticky;top:0;z-index:1}main{max-width:1080px;margin:24px auto;padding:0 20px}button,select,input{font:inherit;padding:6px 12px;margin:3px}button{cursor:pointer}article{background:white;border:1px solid #dbe0e6;border-radius:8px;padding:20px;margin:15px 0}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:14px/1.6 ui-monospace,monospace}textarea{width:98%;min-height:90px;font:inherit}label{display:inline-block;margin:8px 20px 8px 0}.source{white-space:pre-wrap;overflow-wrap:anywhere}summary{cursor:pointer;font-weight:bold}.note{color:#576678;font-size:15px}.tag{font-weight:bold;color:#126285}.controls{display:flex;flex-wrap:wrap;align-items:center}#status{font-size:15px}</style>
<header><strong>Paper 1 · 标注者 __WHO__</strong> <span id="status"></span><div class="controls"><button id="export">导出 JSON 备份</button><label style="margin:0">导入 JSON <input id="import" type="file" accept=".json" style="width:190px"></label><select id="filter"><option value="all">全部条目</option><option value="D">D 输入差异</option><option value="E">E 实际修改</option><option value="todo">未完成</option></select></div></header>
<main><p>先读 <a href="详细说明.html" target="_blank">详细中文说明</a>。独立判断；0/U 请写依据。此页面完全离线。进度缓存于本机，退出前请导出 JSON。</p>
<div class="controls"><button id="prev">上一条</button><button id="next">保存本条并下一条</button><label>跳转 <input id="jump" type="number" min="1" max="400" value="1" style="width:70px"></label><button id="go">前往</button></div>
<article><span class="tag" id="title"></span><p id="operation"></p><h3>原文</h3><div class="source" id="source"></div><h3>允许字段</h3><div id="schema"></div><h3>本条三元组</h3><pre id="triple"></pre><details open><summary>输入图谱</summary><pre id="input"></pre></details><details><summary id="outputTitle">输出图谱</summary><pre id="output"></pre></details></article>
<article><label>输入中对应事实有错误？<select id="is_error"><option value="">未填写</option><option value="1">1 有错误</option><option value="0">0 无错误</option><option value="U">U 无法判断</option></select></label><label>指定操作可接受？<select id="repair_acceptable"><option value="">未填写</option><option value="1">1 可接受</option><option value="0">0 不可接受</option><option value="U">U 无法判断</option></select></label><p class="note">只评价本条指定操作。对 add，“有错误”指指定事实确实缺失且需要补入。D 是建议操作，E 是实际操作。图谱其他位置的错误不要混入本条判断。</p><textarea id="notes" placeholder="判断依据；0/U 请说明原因，建议引用原文"></textarea></article><button id="next2">保存本条并下一条</button></main>
<script>
const DATA=__DATA__,WHO='__WHO__',VERSION='2026-09-21-v1',KEY='paper1-human-'+VERSION+'-'+WHO;
const $=id=>document.getElementById(id);let labels={},index=0;
try{labels=JSON.parse(localStorage.getItem(KEY)||'{}')}catch(e){}
const uid=x=>x.task+':'+x.item_id;
function complete(x){const a=labels[uid(x)]||{};return ['1','0','U'].includes(a.is_error)&&['1','0','U'].includes(a.repair_acceptable)}
function save(){const x=DATA[index];labels[uid(x)]={task:x.task,item_id:x.item_id,is_error:$('is_error').value,repair_acceptable:$('repair_acceptable').value,notes:$('notes').value};try{localStorage.setItem(KEY,JSON.stringify(labels))}catch(e){alert('缓存不可用，请立即导出 JSON 备份。')}status()}
function status(){$('status').textContent=` 已完成 ${DATA.filter(complete).length}/400 · 当前 ${index+1}/400`}
function render(){const x=DATA[index],a=labels[uid(x)]||{};$('title').textContent=`${x.task==='D'?'D 输入差异（建议操作）':'E 实际修改'} · ${x.item_id}`;$('operation').textContent=`操作：${x.operation==='add'?'增加 add':'删除 remove'}`;$('source').textContent=x.source;$('schema').textContent=x.schema.join('；');for(const [id,key]of [['triple','triple'],['input','input_graph'],['output','output_graph']])$(id).textContent=x[key]===null?'本任务没有实际修复输出。请评价上方建议操作。':JSON.stringify(x[key],null,2);for(const k of ['is_error','repair_acceptable','notes'])$(k).value=a[k]||'';$('jump').value=index+1;status()}
function move(step){save();const mode=$('filter').value;let next=index;for(let n=0;n<DATA.length;n++){next=(next+step+DATA.length)%DATA.length;if(mode==='all'||DATA[next].task===mode||(mode==='todo'&&!complete(DATA[next]))){index=next;render();window.scrollTo(0,0);return}}alert('当前筛选下没有未完成条目。')}
$('next').onclick=$('next2').onclick=()=>move(1);$('prev').onclick=()=>move(-1);$('go').onclick=()=>{save();index=Math.max(0,Math.min(DATA.length-1,Number($('jump').value)-1));render()};$('filter').onchange=()=>move(1);
for(const k of ['is_error','repair_acceptable','notes'])$(k).onchange=save;
$('export').onclick=()=>{save();const payload={version:VERSION,annotator:WHO,labels:DATA.map(x=>labels[uid(x)]||{task:x.task,item_id:x.item_id,is_error:'',repair_acceptable:'',notes:''})};const a=document.createElement('a'),url=URL.createObjectURL(new Blob([JSON.stringify(payload,null,2)],{type:'application/json'}));a.href=url;a.download='paper1_annotator_'+WHO+'.json';a.click();setTimeout(()=>URL.revokeObjectURL(url),1000)};
$('import').onchange=async e=>{try{const p=JSON.parse(await e.target.files[0].text());if(p.annotator!==WHO||p.version!==VERSION||!Array.isArray(p.labels))throw Error('身份或版本不匹配');const allowed=new Set(DATA.map(uid)),next={};for(const x of p.labels){if(!allowed.has(uid(x))||next[uid(x)])throw Error('未知或重复 ID');for(const k of ['is_error','repair_acceptable'])if(!['','0','1','U'].includes(x[k]))throw Error('无效标签');next[uid(x)]=x}if(!confirm('导入会替换当前页面进度。已导出备份吗？'))return;labels=next;localStorage.setItem(KEY,JSON.stringify(labels));render()}catch(err){alert('导入失败：'+err.message)}};
window.addEventListener('beforeunload',save);render();
</script></html>'''


def main():
    records, private = items()
    OUT.mkdir(exist_ok=True)
    for who in ['A','B']:
        folder=OUT/f'annotator_{who}';folder.mkdir(exist_ok=True)
        (folder/'ANNOTATOR_GUIDE_zh.md').write_text((HERE/'ANNOTATOR_GUIDE_zh.md').read_text())
        guide=markdown.markdown((HERE/'ANNOTATOR_GUIDE_zh.md').read_text(),extensions=['tables','fenced_code'])
        (folder/'详细说明.html').write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>真人标注详细说明</title><style>body{max-width:980px;margin:40px auto;padding:20px;font:17px/1.8 system-ui,sans-serif;color:#172536}table{border-collapse:collapse}th,td{border:1px solid #ccd;padding:8px}pre{white-space:pre-wrap}code{background:#f3f5f7}</style>'+guide+'</html>')
        data=json.dumps(records,ensure_ascii=False).replace('<','\\u003c')
        (folder/'标注页面.html').write_text(HTML.replace('__WHO__',who).replace('__DATA__',data))
        workbook(records,folder/'标注表.xlsx',who)
        with zipfile.ZipFile(OUT/f'标注者_{who}.zip','w',zipfile.ZIP_DEFLATED) as z:
            for p in sorted(folder.iterdir()):z.write(p,p.name)
    coord=OUT/'coordinator';coord.mkdir(exist_ok=True)
    (coord/'private_mapping.json').write_text(json.dumps(private,ensure_ascii=False,indent=2)+'\n')
    (coord/'COORDINATOR_GUIDE_zh.md').write_text((HERE/'COORDINATOR_GUIDE_zh.md').read_text())
    manifest={'version':'2026-09-21-v1','tasks':{'D':200,'E':200},'labels':'blank; actual human review pending',
        'frozen_inputs':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [EXT/'natural_error_annotation_sample.csv',AUDIT/'annotator_a.csv',AUDIT/'annotator_b.csv',AUDIT/'manifest.json']},
        'public_items_sha256':hashlib.sha256(json.dumps(records,ensure_ascii=False,sort_keys=True).encode()).hexdigest(),
        'deliveries':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [OUT/'标注者_A.zip', OUT/'标注者_B.zip']}}
    (OUT/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    with zipfile.ZipFile(OUT/'协调员.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(coord.iterdir()):z.write(p,p.name)
        z.write(OUT/'manifest.json','manifest.json')
    print('Built 400-item independent packages A/B; original frozen samples unchanged.')


if __name__=='__main__':main()
